"""Génération du modèle Excel officiel UCA."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.urbanism_import.sheet_registry import SHEET_DEFINITIONS

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def generate_official_template() -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for defn in SHEET_DEFINITIONS:
        ws = wb.create_sheet(defn.sheet_key)
        for col_idx, header in enumerate(defn.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            ws.column_dimensions[cell.column_letter].width = max(14, len(header) + 2)
        _add_sample_row(ws, defn.sheet_key)

    readme = wb.create_sheet("README", 0)
    readme["A1"] = "Urban Cyber Architect — Modèle cartographie urbanisme SI V1.4"
    readme["A1"].font = Font(bold=True, size=14)
    readme["A3"] = "Remplissez les feuilles 01 à 15. Les ID doivent être uniques par type."
    readme["A4"] = "Les colonnes de référence (Métier associé, Objectif, etc.) utilisent l'ID ou le Nom."
    readme["A5"] = "Import : module Urbanisme → Importer une cartographie."

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_sample_row(ws, sheet_key: str) -> None:
    samples = {
        "01_Metiers": ("M1", "Sécurité publique", "Pilotage sécurité", "RSSI", "élevée"),
        "02_Objectifs": ("O1", "Réduire les incidents", "M1", "Objectif stratégique"),
        "03_Processus": ("P1", "Gestion des incidents", "O1", "Processus SOC"),
        "04_Activites": ("A1", "Détection", "P1", "Surveillance SI"),
        "05_Classes": ("C1", "Alerte", "A1", "Classe d'alerte"),
        "06_Organisation": ("ORG1", "DSI", "SOC", "Chef SOC", "Direction sécurité"),
        "07_Operations": ("OP1", "Analyse alertes", "ORG1", "Opération SOC"),
        "08_Fonctionnel": ("F1", "Supervision SOC", "C1", "Fonction supervision"),
        "09_Applicatif": ("APP1", "SIEM", "F1", "Elastic", "8.12", "critique"),
        "10_Technique": ("T1", "Déploiement SIEM", "APP1", "SRV1", "NET1", "SITE1", "Non", "Prod"),
        "11_Serveurs": ("SRV1", "srv-siem-01", "Linux", "10.0.1.10", "VM", "critique"),
        "12_Reseaux": ("NET1", "LAN-SOC", "100", "10.0.1.0/24", "SITE1"),
        "13_Sites": ("SITE1", "Datacenter Paris", "Paris", "1 rue Example"),
        "14_Equipements": ("EQ1", "Poste analyste", "Poste", "Dell", "SRV1", "NET1"),
        "15_Flux": ("APP1", "APP2", "API", "HTTPS", "Échange tickets"),
    }
    row = samples.get(sheet_key)
    if row:
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=2, column=col_idx, value=value)
    if sheet_key == "09_Applicatif":
        ws.cell(row=3, column=1, value="APP2")
        ws.cell(row=3, column=2, value="Ticketing")
        ws.cell(row=3, column=3, value="F1")
        ws.cell(row=3, column=4, value="ServiceNow")
        ws.cell(row=3, column=5, value="Yokohama")
        ws.cell(row=3, column=6, value="moyenne")
