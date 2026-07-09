"""Lecteur Excel (.xlsx) — feuilles officielles UCA."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.urbanism_import.readers.base import ImportReader
from app.services.urbanism_import.readers.row_parser import parse_sheet_rows
from app.services.urbanism_import.sheet_registry import SHEET_DEFINITIONS, resolve_sheet_name
from app.services.urbanism_import.types import ImportCounts, ImportIssue, ImportMode, ParsedImport


class ExcelImportReader(ImportReader):
    def read(self, content: bytes, filename: str) -> ParsedImport:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        parsed = ParsedImport()
        for sheet_name in wb.sheetnames:
            resolved = resolve_sheet_name(sheet_name)
            if not resolved:
                continue
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(c).strip() if c is not None else "" for c in rows[0]]
            data_rows: list[dict[str, Any]] = []
            for idx, row in enumerate(rows[1:], start=2):
                if not any(cell is not None and str(cell).strip() for cell in row):
                    continue
                cells = {
                    headers[i]: row[i] if i < len(row) else None
                    for i in range(len(headers))
                    if headers[i]
                }
                data_rows.append({"row_number": idx, "cells": cells})
            sheet_rows, sheet_flux, sheet_issues = parse_sheet_rows(resolved, data_rows)
            parsed.rows.extend(sheet_rows)
            parsed.flux_rows.extend(sheet_flux)
            parsed.issues.extend(sheet_issues)
        wb.close()
        parsed.counts = _compute_counts(parsed)
        return parsed


def _compute_counts(parsed: ParsedImport) -> ImportCounts:
    counts = ImportCounts()
    for row in parsed.rows:
        key = {
            "metier": "metiers",
            "objectif": "objectifs",
            "processus": "processus",
            "activite": "activites",
            "classe": "classes",
            "organisation": "organisations",
            "operation": "operations",
            "ilot_fonctionnel": "fonctions",
            "ilot_applicatif": "applications",
            "serveur": "serveurs",
            "reseau": "reseaux",
            "site": "sites",
            "poste_travail": "equipements",
        }.get(row.entity_type)
        if key:
            setattr(counts, key, getattr(counts, key) + 1)
    counts.flux = len(parsed.flux_rows)
    return counts


def empty_workbook_template_bytes() -> bytes:
    from app.services.urbanism_import.template_generator import generate_official_template

    return generate_official_template()
