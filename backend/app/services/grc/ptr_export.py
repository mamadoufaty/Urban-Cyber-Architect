"""Export PTR — CSV, Excel, PDF."""

from __future__ import annotations

import csv
import io
from typing import Any

from app.services.grc.ptr_service import EXPORT_HEADERS


def export_ptr_csv(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow([label for _, label in EXPORT_HEADERS])
    for row in rows:
        writer.writerow([row.get(key, "") for key, _ in EXPORT_HEADERS])
    return buffer.getvalue().encode("utf-8-sig")


def export_ptr_xlsx(rows: list[dict[str, Any]], *, summary: dict[str, Any] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "PTR"

    if summary:
        ws.append(["Plan de Traitement des Risques"])
        ws.append(["Projet", summary.get("project_name", "")])
        ws.append(["Actions totales", summary.get("total_actions", 0)])
        ws.append(["Avancement global (%)", summary.get("global_progress_percent", 0)])
        ws.append([])

    headers = [label for _, label in EXPORT_HEADERS]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(key, "") for key, _ in EXPORT_HEADERS])

    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max_len + 2, 48)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_ptr_pdf(
    rows: list[dict[str, Any]],
    *,
    summary: dict[str, Any] | None = None,
    project_name: str = "",
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Plan de Traitement des Risques (PTR)", styles["Title"]),
        Paragraph(
            f"Projet : {project_name or (summary or {}).get('project_name', '—')}",
            styles["Normal"],
        ),
    ]
    if summary:
        story.append(
            Paragraph(
                f"Actions : {summary.get('total_actions', 0)} — "
                f"Avancement : {summary.get('global_progress_percent', 0)}%",
                styles["Normal"],
            )
        )
    story.append(Spacer(1, 12))

    compact_headers = [
        ("ptr_id", "ID"),
        ("security_measure", "Mesure"),
        ("responsible", "Responsable"),
        ("priority", "Priorité"),
        ("due_date", "Échéance"),
        ("status", "Statut"),
        ("progress_percent", "%"),
        ("residual_risk", "Risque rés."),
    ]
    table_data = [[label for _, label in compact_headers]]
    for row in rows:
        table_data.append([str(row.get(key, ""))[:80] for key, _ in compact_headers])

    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f8")]),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()
