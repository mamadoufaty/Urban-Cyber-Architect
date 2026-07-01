"""Export SoA — CSV, Excel, PDF."""

from __future__ import annotations

import csv
import io
from typing import Any

EXPORT_HEADERS: list[tuple[str, str]] = [
    ("iso_reference", "Référence ISO"),
    ("control_name", "Nom du contrôle"),
    ("applicable", "Applicable"),
    ("justification", "Justification"),
    ("implemented", "Implémenté"),
    ("ebios_source", "Source EBIOS"),
    ("associated_measure", "Mesure associée"),
    ("decision", "Décision"),
    ("responsible", "Responsable"),
    ("status", "Statut"),
    ("comment", "Commentaire"),
]


def export_soa_csv(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow([label for _, label in EXPORT_HEADERS])
    for row in rows:
        writer.writerow([row.get(key, "") for key, _ in EXPORT_HEADERS])
    return buffer.getvalue().encode("utf-8-sig")


def export_soa_xlsx(rows: list[dict[str, Any]], *, summary: dict[str, Any] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "SoA"

    if summary:
        ws.append(["Déclaration d'Applicabilité"])
        ws.append(["Version", summary.get("soa_version", "")])
        ws.append(["Projet", summary.get("project_name", "")])
        ws.append(["Généré le", str(summary.get("generated_at", ""))])
        ws.append([])

    headers = [label for _, label in EXPORT_HEADERS]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row.get(key, "") for key, _ in EXPORT_HEADERS])

    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max_len + 2, 48)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_soa_pdf(
    rows: list[dict[str, Any]],
    *,
    summary: dict[str, Any] | None = None,
    project_name: str = "",
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Déclaration d'Applicabilité (SoA)", styles["Title"]),
        Paragraph(
            f"Projet : {project_name or summary.get('project_name', '—') if summary else '—'} — lecture seule",
            styles["Normal"],
        ),
    ]
    if summary:
        story.append(
            Paragraph(
                f"Version : {summary.get('soa_version', '')} — "
                f"Couverture : {summary.get('coverage_rate_percent', 0)}%",
                styles["Normal"],
            )
        )
    story.append(Spacer(1, 12))

    headers = [label for _, label in EXPORT_HEADERS]
    table_data = [headers]
    for row in rows:
        table_data.append([str(row.get(key, ""))[:120] for key, _ in EXPORT_HEADERS])

    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f8")]),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()
